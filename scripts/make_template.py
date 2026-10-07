"""Build templates/MLP_Register_Template.xlsx (the blank register each staff member fills in).

  python scripts/make_template.py

Columns must match what scripts/import_register.py reads. Row 1 = headers (* = required),
row 2 = example (skipped on import), real data from row 3.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "MLP_Register_Template.xlsx"
ROWS = 500                                   # rows covered by dropdowns / text formatting

# sheet -> [(column, required, example, note, dropdown values or None, text format?)]
SHEETS = {
    "Staff": [
        ("staff_name", True, "Nash Lim", None, None, False),
        ("email", True, "nash@example.com", "Used to log in to the portal", None, False),
        ("phone", False, "0123456789", None, None, True),
        ("role", True, "staff", "admin sees everything; viewer is read-only", ["admin", "staff", "viewer"], False),
    ],
    "Landlords": [
        ("landlord_code", True, "LL-TAN01", "Short and permanent", None, False),
        ("name", True, "Tan Ah Kow", None, None, False),
        ("phone", False, "0198765432", None, None, True),
        ("email", False, "tan@example.com", None, None, False),
        ("management_fee_pct", False, "10", "Defaults to 10. No bank details here: the admin enters them in the portal", None, False),
    ],
    "Units": [
        ("unit_code", True, "SNDN-3418", "Area + lot/unit no. Short, unique, permanent: used everywhere incl. bank matching", None, False),
        ("address", True, "Lot 3418, Taman Senadin, 98100 Miri", None, None, False),
        ("area", False, "Senadin", None, None, False),
        ("unit_type", False, "Terrace house", None, None, False),
        ("landlord_code", True, "LL-TAN01", "Must exist in the Landlords sheet", None, False),
        ("assigned_staff_email", True, "nash@example.com", "Must exist in the Staff sheet", None, False),
        ("status", True, "occupied", None, ["occupied", "vacant", "inactive"], False),
        ("notes", False, "", None, None, False),
    ],
    "Tenants": [
        ("tenant_name", True, "Ahmad bin Ali", None, None, False),
        ("phone", True, "0123456789", "As 0123456789 (WhatsApp number). No IC numbers anywhere", None, True),
        ("email", False, "", None, None, False),
        ("notes", False, "", None, None, False),
    ],
    "Tenancies": [
        ("unit_code", True, "SNDN-3418", None, None, False),
        ("tenant_phone", True, "0123456789", "Must exist in the Tenants sheet", None, True),
        ("start_date", True, "01/03/2026", "DD/MM/YYYY", None, True),
        ("end_date", False, "28/02/2027", "DD/MM/YYYY, blank if open-ended", None, True),
        ("monthly_rent", True, "1200", None, None, False),
        ("due_day", True, "1", "Day of the month rent is due, 1-28", None, False),
        ("deposit_rental", False, "2400", None, None, False),
        ("deposit_utility", False, "600", None, None, False),
        ("status", True, "active", "One active tenancy per unit", ["upcoming", "active", "ended"], False),
    ],
    "Utility_Accounts": [
        ("unit_code", True, "SNDN-3418", None, None, False),
        ("type", True, "electric", None, ["electric", "water"], False),
        ("account_no", True, "220012345678", "Enter as text so leading zeros stay", None, True),
    ],
    "Payment_History": [
        ("unit_code", True, "SNDN-3418", "Optional sheet: last 3-6 months of payments", None, False),
        ("period", True, "2026-09", "Rent month as YYYY-MM", None, True),
        ("amount_paid", True, "1200", None, None, False),
        ("paid_date", True, "02/09/2026", "DD/MM/YYYY", None, True),
        ("method", False, "bank_transfer", None, ["bank_transfer", "duitnow", "cash", "cheque", "other"], False),
        ("reference", False, "MBB123456", None, None, False),
    ],
}

INSTRUCTIONS = """MLP Rental Register: how to fill in

1. Fill in the sheets in this order: Staff, Landlords, Units, Tenants, Tenancies, Utility_Accounts, Payment_History.
2. Row 2 of every sheet is an EXAMPLE (grey) and is skipped on import. Enter real data from row 3.
3. Columns marked * are required.
4. Dates: DD/MM/YYYY (or real Excel dates). Payment_History period: YYYY-MM.
5. Phones: 0123456789 format. They are converted to 60123456789 for WhatsApp.
6. Unit codes: short, unique and permanent, e.g. SNDN-3418 (area + lot/unit number).
7. One ACTIVE tenancy per unit. Use 'upcoming' for a signed tenancy that hasn't started.
8. Do NOT enter IC numbers or landlord bank details. The admin enters bank details in the portal.
9. Don't rename sheets or header cells; the importer reads them by name.

Check the file:   python scripts/import_register.py MLP_Register.xlsx
Save to database: python scripts/import_register.py MLP_Register.xlsx --commit
"""

HEAD = PatternFill("solid", fgColor="1F4E78")
REQ = PatternFill("solid", fgColor="C00000")
EXAMPLE = PatternFill("solid", fgColor="E7E6E6")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    for i, line in enumerate(INSTRUCTIONS.splitlines(), start=1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 120

    for sheet, cols in SHEETS.items():
        ws = wb.create_sheet(sheet)
        ws.freeze_panes = "A3"
        for c, (name, required, example, note, choices, as_text) in enumerate(cols, start=1):
            letter = ws.cell(row=1, column=c).column_letter
            head = ws.cell(row=1, column=c, value=name + ("*" if required else ""))
            head.font = Font(bold=True, color="FFFFFF")
            head.fill = REQ if required else HEAD
            head.alignment = Alignment(horizontal="center")
            if note:
                head.comment = Comment(note, "MLP")
            ex = ws.cell(row=2, column=c, value=example)
            ex.fill = EXAMPLE
            ex.font = Font(italic=True, color="7F7F7F")
            if as_text:                      # keep leading zeros / stop Excel turning values into numbers
                for r in range(2, ROWS + 3):
                    ws.cell(row=r, column=c).number_format = "@"
            if choices:
                dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True,
                                    showErrorMessage=True, errorTitle="Not allowed",
                                    error="Pick one of: " + ", ".join(choices))
                dv.add(f"{letter}3:{letter}{ROWS + 2}")
                ws.add_data_validation(dv)
            ws.column_dimensions[letter].width = max(14, len(name) + 4, len(str(example)) + 2)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
