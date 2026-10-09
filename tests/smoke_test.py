"""End-to-end checks (Setup Guide §10, the parts that don't need a browser) on a THROWAWAY database.

  MLP_TEST_DATABASE_URL=postgresql://postgres@localhost:55432/mlp_test python tests/smoke_test.py

WARNING: drops and recreates the public schema of that database. Never point it at Supabase.
"""
import datetime as dt
import io
import os
import subprocess
import sys
import tempfile
from pathlib import Path

import psycopg
from openpyxl import load_workbook
from psycopg.rows import dict_row

ROOT = Path(__file__).resolve().parent.parent
URL = os.environ.get("MLP_TEST_DATABASE_URL")
if not URL or "supabase" in URL:
    sys.exit(__doc__)
os.environ["DATABASE_URL"] = URL
sys.path[:0] = [str(ROOT / "worker"), str(ROOT / "portal")]

TODAY = dt.date.today()
THIS_MONTH = TODAY.replace(day=1)
LAST_MONTH = (THIS_MONTH - dt.timedelta(days=1)).replace(day=1)
MID_START = THIS_MONTH.replace(day=3)          # mid-month start → pro-rata note
passed = 0


def check(cond, msg):
    global passed
    if not cond:
        raise AssertionError(msg)
    passed += 1
    print(f"  ✓ {msg}")


def db():
    return psycopg.connect(URL, row_factory=dict_row, autocommit=True)


def one(sql, params=None):
    with db() as c:
        return c.execute(sql, params).fetchone()


def reset_schema():
    with db() as c:
        c.execute("DROP SCHEMA public CASCADE; CREATE SCHEMA public;")
        for f in sorted((ROOT / "db/migrations").glob("*.sql")):
            c.execute(f.read_text())


def workbook(path, units_landlord="LL-TAN01", second_active=False):
    wb = load_workbook(ROOT / "templates/MLP_Register_Template.xlsx")
    data = {
        "Staff": [["Nash Lim", "nash@mlp.test", "0120000001", "admin"],
                  ["Mei Ling", "Mei@MLP.test", "0120000002", "staff"]],
        "Landlords": [["LL-TAN01", "Tan Ah Kow", "0198765432", "", "10"],
                      ["ll-wong", "Wong Siew", "", "", ""]],
        "Units": [["SNDN-3418", "Lot 3418, Senadin", "Senadin", "Terrace", units_landlord, "mei@mlp.test", "occupied", ""],
                  ["PJYA-12", "12 Pujut", "Pujut", "Apartment", "LL-WONG", "nash@mlp.test", "occupied", ""],
                  ["LUTG-7", "7 Lutong", "Lutong", "", "LL-WONG", "mei@mlp.test", "vacant", ""]],
        "Tenants": [["Ahmad bin Ali", "012-345 6789", "", ""],
                    ["Siti Aminah", "+60 19-888 7777", "", ""],
                    ["Lee Chong", "0111222333", "", ""]],
        "Tenancies": [["SNDN-3418", "0123456789", "01/03/2026", "", "1200", "1", "2400", "600", "active"],
                      # a real Excel date cell (not text) to exercise the date fix
                      ["PJYA-12", "0198887777", dt.datetime.combine(MID_START, dt.time()), "", "950", "5", "", "", "active"],
                      ["LUTG-7", "0111222333", "01/01/2027", "", "800", "1", "", "", "upcoming"]]
                     + ([["SNDN-3418", "0111222333", "01/06/2026", "", "1000", "1", "", "", "active"]]
                        if second_active else []),
        "Utility_Accounts": [["SNDN-3418", "electric", "220012345678"],
                             ["PJYA-12", "electric", "220099999999"]],
        "Payment_History": [["SNDN-3418", f"{LAST_MONTH:%Y-%m}", "1200", f"02/{LAST_MONTH:%m/%Y}", "duitnow", "OLD1"]],
    }
    for sheet, rows in data.items():
        ws = wb[sheet]
        for r, row in enumerate(rows, start=3):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)
    wb.save(path)


def importer(path, commit=False):
    args = [sys.executable, str(ROOT / "scripts/import_register.py"), str(path)] + (["--commit"] if commit else [])
    return subprocess.run(args, capture_output=True, text=True, env={**os.environ, "DATABASE_URL": URL}).stdout


def counts():
    return {t: one(f"SELECT count(*) AS n FROM {t}")["n"]
            for t in ("staff", "landlords", "units", "tenants", "tenancies", "utility_accounts",
                      "payments", "rent_schedule")}


def main():
    tmp = Path(tempfile.mkdtemp())
    reset_schema()
    print("Register import")
    workbook(tmp / "bad.xlsx", units_landlord="LL-NOPE")
    out = importer(tmp / "bad.xlsx", commit=True)
    check("Units row 3: unknown landlord_code 'LL-NOPE'" in out, "1. bad landlord code reported by sheet/row")
    check(counts()["staff"] == 0, "1. nothing saved when there are problems")

    workbook(tmp / "dup.xlsx", second_active=True)
    out = importer(tmp / "dup.xlsx")
    check("unit already has an active tenancy" in out, "3. two active tenancies for one unit rejected")

    workbook(tmp / "good.xlsx")
    out = importer(tmp / "good.xlsx")
    check("Validation OK" in out and counts()["staff"] == 0, "validate-only run saves nothing")
    out = importer(tmp / "good.xlsx", commit=True)
    check("Imported and committed" in out, f"import committed ({out.strip().splitlines()[-1]})")
    first = counts()
    importer(tmp / "good.xlsx", commit=True)
    check(counts() == first, "2. re-import creates no duplicates")
    check(one("SELECT phone FROM tenants WHERE name = 'Siti Aminah'")["phone"] == "60198887777",
          "phones normalised to 60…")
    check(one("SELECT start_date FROM tenancies t JOIN units u ON u.id = t.unit_id WHERE u.code = 'PJYA-12'")
          ["start_date"] == MID_START, "real Excel date cell imported as the right day (not day/month swapped)")
    check(one("SELECT status FROM units WHERE code = 'LUTG-7'")["status"] == "vacant",
          "unit with only an upcoming tenancy stays vacant")

    print("Rent schedule")
    n1 = one("SELECT generate_rent_schedule(%s) AS n", (THIS_MONTH,))["n"]
    n2 = one("SELECT generate_rent_schedule(%s) AS n", (THIS_MONTH,))["n"]
    check(n1 == 2 and n2 == 0, f"4. generate twice: {n1} then {n2} lines")
    check(one("SELECT note FROM v_rent_status WHERE unit_code = 'PJYA-12' AND period = %s",
              (THIS_MONTH,))["note"] == "First month: check pro-rata", "5. mid-month start flagged pro-rata")
    check(one("SELECT due_date FROM v_rent_status WHERE unit_code = 'PJYA-12' AND period = %s",
              (THIS_MONTH,))["due_date"] == MID_START.replace(day=5), "due date = due_day, never before start")
    check(one("SELECT status FROM v_rent_status WHERE unit_code = 'SNDN-3418' AND period = %s",
              (LAST_MONTH,))["status"] == "PAID", "payment history imported as PAID")

    print("Payments")
    with db() as c:
        s = c.execute("SELECT schedule_id, tenancy_id FROM v_rent_status WHERE unit_code = 'PJYA-12' "
                      "AND period = %s", (THIS_MONTH,)).fetchone()
        c.execute("INSERT INTO payments (tenancy_id, rent_schedule_id, amount, paid_date) VALUES (%s, %s, 500, %s)",
                  (s["tenancy_id"], s["schedule_id"], TODAY))
        c.execute("UPDATE rent_schedule SET due_date = %s WHERE id = %s",
                  (TODAY + dt.timedelta(days=3), s["schedule_id"]))
        r = c.execute("SELECT status, balance FROM v_rent_status WHERE schedule_id = %s", (s["schedule_id"],)).fetchone()
        check(r["status"] == "PARTIAL" and r["balance"] == 450, "8. partial payment → PARTIAL, balance 450")
        c.execute("UPDATE rent_schedule SET due_date = %s WHERE id = %s",
                  (TODAY - dt.timedelta(days=4), s["schedule_id"]))
        r = c.execute("SELECT status, days_late, grace_end FROM v_rent_status WHERE schedule_id = %s",
                      (s["schedule_id"],)).fetchone()
        check(r["status"] == "GRACE" and r["days_late"] == 4 and r["grace_end"] == TODAY + dt.timedelta(days=3),
              "8. 4 days past due → in the 7-day grace period")
        c.execute("UPDATE rent_schedule SET due_date = %s WHERE id = %s",
                  (TODAY - dt.timedelta(days=8), s["schedule_id"]))
        r = c.execute("SELECT status, days_late FROM v_rent_status WHERE schedule_id = %s", (s["schedule_id"],)).fetchone()
        check(r["status"] == "OVERDUE" and r["days_late"] == 8, "8. after the grace period → OVERDUE, 8 days late")
        pid = c.execute("INSERT INTO payments (tenancy_id, rent_schedule_id, amount, paid_date) "
                        "VALUES (%s, %s, 450, %s) RETURNING id",
                        (s["tenancy_id"], s["schedule_id"], TODAY)).fetchone()["id"]
        check(c.execute("SELECT status FROM v_rent_status WHERE schedule_id = %s",
                        (s["schedule_id"],)).fetchone()["status"] == "PAID", "7. fully paid → PAID")
        c.execute("UPDATE payments SET voided = TRUE, void_reason = 'test' WHERE id = %s", (pid,))
        r = c.execute("SELECT status, paid FROM v_rent_status WHERE schedule_id = %s", (s["schedule_id"],)).fetchone()
        check(r["status"] == "OVERDUE" and r["paid"] == 500, "10. voided payment excluded from totals")

    print("Bank matching")
    from bank import BANK_FORMATS, load_statement, row_hash, strong
    fmt = BANK_FORMATS["Company Bank (generic CSV)"]
    stmt = load_statement(open(ROOT / "tests/fixtures/sample_bank.csv", "rb"), fmt)
    check(stmt["amount"].tolist() == [1200.0, 950.0, 300.0], "CSV: credits only, '1,200.00' parsed")
    signed = io.BytesIO(b"Txn Date,Details,Amount\n01/10/2026,RENT A,1200.00\n02/10/2026,FEE,-5.00\n")
    s2 = load_statement(signed, {"skiprows": 0, "date": "Txn Date", "desc": "Details", "ref": None,
                                 "amount": "Amount", "date_format": "%d/%m/%Y"})
    check(s2["amount"].tolist() == [1200.0], "CSV: signed Amount column")
    drcr = io.BytesIO(b"Bank export\nDate,Desc,Amt,Type\n2026-10-01,RENT B,950.00,CR\n2026-10-02,X,10.00,DR\n")
    s3 = load_statement(drcr, {"skiprows": 1, "date": "Date", "desc": "Desc", "ref": None, "amount": "Amt",
                               "drcr": "Type", "credit_marker": "CR", "date_format": None})
    check(s3["amount"].tolist() == [950.0] and str(s3["txn_date"][0]) == "2026-10-01",
          "CSV: DR/CR column + header line skipped")
    with db() as c:
        def upload():
            return sum(c.execute("""INSERT INTO bank_transactions (account_label, txn_date, description,
                                        reference, amount, row_hash) VALUES (%s, %s, %s, %s, %s, %s)
                                    ON CONFLICT (row_hash) DO NOTHING""",
                                 ("Company Bank", r.txn_date, r.description, r.reference or None,
                                  float(r.amount), row_hash("Company Bank", r))).rowcount
                       for r in stmt.itertuples())
        a, b = upload(), upload()
        check(a == 3 and b == 0, f"13. bank CSV uploaded twice: {a} then {b} rows")
    check(strong("DUITNOW TRF AHMAD BIN ALI SNDN3418", "SNDN-3418", "Ahmad bin Ali"), "unit code w/o dash ranks strong")
    check(not strong("CASH DEPOSIT", "SNDN-3418", "Ahmad bin Ali"), "unrelated text is not strong")

    print("Messages")
    from messages import norm_phone, wa_link
    link = wa_link("60123456789", "Hi Ahmad,\nRM1,200.00 due 01/10 & thanks")
    check(link == "https://wa.me/60123456789?text=Hi%20Ahmad%2C%0ARM1%2C200.00%20due%2001/10%20%26%20thanks",
          "12a. wa.me link keeps line breaks, RM amounts and &")
    check(norm_phone("+60 12-345 6789") == norm_phone("012-345 6789") == "60123456789", "portal norm_phone")

    print("Worker")
    os.chdir(tmp)                                    # jobs write raw/ and state/ relative to cwd
    from jobs import electric, rent
    rent.generate_this_month()
    r = one("SELECT status, rows_affected FROM job_runs ORDER BY id DESC LIMIT 1")
    check(r["status"] == "ok" and r["rows_affected"] == 0, "rent_this_month job logged ok (0 new lines)")

    def fake_fetch(day_dir, **kw):                  # SEBCares JSON as the bill page receives it
        import json
        day_dir.mkdir(parents=True, exist_ok=True)
        bill = {"ContractSubscription": {"Id": "1", "ContractId": "220012345678", "Nickname": "SNDN"},
                "CustomerInformations": {
                    "T_ACCOUNT_BALANCES": {"item": {"List": [{"TEXT": "Open", "WITHD_VAL": "85.2000"}]}},
                    "T_INVOICE_LIST": {"item": {"List": [{"PRINT_DOC": "7360001", "INVOICE_DATE": "2026-09-29",
                                                          "DUE_DATE": "2026-10-20", "CURR_AMT": "     85.20",
                                                          "CURR_TOTAL_AMT": "     85.20"}]}}}}
        (day_dir / "220012345678.json").write_text(json.dumps(bill))
        return [], ["220099999999: bill data did not load"]

    electric.fetch = fake_fetch
    os.environ["SEB_USERNAME"] = "test@example.invalid"
    electric.run()
    r = one("SELECT status, rows_affected, error FROM job_runs WHERE job = 'electric_check' ORDER BY id DESC LIMIT 1")
    saved = one("SELECT count(*) AS n FROM bill_checks")["n"]
    check(saved == 1 and r["rows_affected"] == 1, "15. good account saved although one failed")
    check(r["status"] == "failed" and "220099999999" in r["error"], "15. failure + account listed in job_runs")
    check(float(one("SELECT outstanding FROM v_electric_latest WHERE unit_code = 'SNDN-3418'")["outstanding"]) == 85.20,
          "v_electric_latest shows the parsed bill")
    check(one("SELECT due_date FROM utility_bills WHERE bill_no = '7360001'")["due_date"] == dt.date(2026, 10, 20),
          "SEB bill kept in utility_bills")

    print(f"\nAll {passed} checks passed.")


if __name__ == "__main__":
    main()
